# src/rag/loader.py
from pathlib import Path
from dataclasses import dataclass, field
from loguru import logger
import pypdf
import docx


@dataclass
class Document:
    """A single unit of loaded text with its source metadata."""
    content: str
    source: str
    metadata: dict = field(default_factory=dict)


def _load_pdf(path: Path) -> str:
    reader = pypdf.PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _load_docx(path: Path) -> str:
    doc = docx.Document(str(path))
    return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])


def _load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


LOADERS = {
    ".pdf":  _load_pdf,
    ".docx": _load_docx,
    ".txt":  _load_txt,
    ".md":   _load_txt,
}


def load_documents(input_dir: str) -> list[Document]:
    """
    Recursively load all supported files from a directory.
    Returns a list of Document objects.
    """
    input_path = Path(input_dir)
    if not input_path.exists():
        raise FileNotFoundError(f"Input directory not found: {input_dir}")

    documents = []

    for file_path in sorted(input_path.rglob("*")):
        suffix = file_path.suffix.lower()
        if suffix not in LOADERS:
            continue

        try:
            logger.info(f"Loading: {file_path.name}")
            content = LOADERS[suffix](file_path)

            if not content.strip():
                logger.warning(f"Empty content, skipping: {file_path.name}")
                continue

            documents.append(Document(
                content=content,
                source=str(file_path),
                metadata={"filename": file_path.name, "type": suffix.lstrip(".")},
            ))

        except Exception as e:
            logger.error(f"Failed to load {file_path.name}: {e}")
            continue

    logger.success(f"Loaded {len(documents)} documents from {input_dir}")
    return documents